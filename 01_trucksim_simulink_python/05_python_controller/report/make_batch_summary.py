#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成一份批量仿真总报告。

输入为 run_batch_python.m 写出的 batch_run_dirs.txt。输出目录中包含：
  figures/                 各工况曲线图
  批量指标汇总.csv
  阶段二批量仿真测试报告.docx
"""
import argparse
import csv
import os
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor

from make_report_python import (FIG_SPECS, add_heading, add_run, build_table,
                                load_run, metrics, plot_figs, write_metrics_csv)


def _fmt(value):
    try:
        return "%.4g" % float(value)
    except (TypeError, ValueError):
        return ""


def _add_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, title, size=Pt(20), bold=True, color=RGBColor(0x1F, 0x3F, 0x76))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, subtitle, size=Pt(10.5), color=RGBColor(0x5B, 0x65, 0x73))


def _add_summary_table(doc, results):
    add_heading(doc, "1. 批量概况")
    tbl = doc.add_table(rows=1 + len(results), cols=10)
    tbl.style = "Table Grid"
    header = ["序号", "运行编号", "工况名称", "输入", "初始车速", "结束车速",
              "最大|β|", "最大横摆", "最大d1", "最大d3"]
    for j, htxt in enumerate(header):
        cell = tbl.rows[0].cells[j]
        cell.paragraphs[0].text = ""
        add_run(cell.paragraphs[0], htxt, size=Pt(8), bold=True)

    for i, item in enumerate(results, start=1):
        run_no, case, m, _fig_dir = item
        vals = [
            str(i),
            run_no,
            case.get("case_name", ""),
            case.get("steer_input_type", ""),
            _fmt(m["v_start_kmh"]) + " km/h",
            _fmt(m["v_end_kmh"]) + " km/h",
            _fmt(m["max_abs_beta_deg"]) + " deg",
            _fmt(m["max_abs_yaw_degps"]) + " deg/s",
            _fmt(m["max_abs_delta1_deg"]) + " deg",
            _fmt(m["max_abs_delta3_deg"]) + " deg",
        ]
        for j, val in enumerate(vals):
            cell = tbl.rows[i].cells[j]
            cell.paragraphs[0].text = ""
            add_run(cell.paragraphs[0], val, size=Pt(8))


def _add_case_table(doc, case, m):
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    header = ["参数", "数值", "指标", "数值"]
    for j, htxt in enumerate(header):
        cell = tbl.rows[0].cells[j]
        cell.paragraphs[0].text = ""
        add_run(cell.paragraphs[0], htxt, size=Pt(8.5), bold=True)

    rows = [
        ("工况说明", case.get("description", ""), "起始车速", _fmt(m["v_start_kmh"]) + " km/h"),
        ("输入类型", case.get("steer_input_type", ""), "结束车速", _fmt(m["v_end_kmh"]) + " km/h"),
        ("仿真时长", case.get("stop_time_s", "") + " s", "最大车速", _fmt(m["v_max_kmh"]) + " km/h"),
        ("初始车速", case.get("initial_speed_kmh", "") + " km/h", "最大|β|", _fmt(m["max_abs_beta_deg"]) + " deg"),
        ("目标车速", case.get("target_speed_kmh", case.get("initial_speed_kmh", "")) + " km/h",
         "最大横摆", _fmt(m["max_abs_yaw_degps"]) + " deg/s"),
        ("转向幅值", case.get("steer_amplitude_deg", case.get("steer_step_deg", "-")) + " deg",
         "最大d1", _fmt(m["max_abs_delta1_deg"]) + " deg"),
        ("转向频率/开始", "%s Hz / %s s" % (case.get("steer_frequency_hz", "-"),
                                      case.get("steer_start_time_s", "-")),
         "最大d3", _fmt(m["max_abs_delta3_deg"]) + " deg"),
        ("路面附着系数", case.get("road_friction", "-"), "控制周期", case.get("control_dt_s", "0.01") + " s"),
    ]
    for left_k, left_v, right_k, right_v in rows:
        cells = tbl.add_row().cells
        vals = [left_k, left_v, right_k, right_v]
        for j, val in enumerate(vals):
            cells[j].paragraphs[0].text = ""
            add_run(cells[j].paragraphs[0], val, size=Pt(8.5), bold=(j in (0, 2)))


def _add_case_figures(doc, fig_dir):
    for fname, title, _key, _unit in FIG_SPECS:
        path = os.path.join(fig_dir, fname)
        if not os.path.isfile(path):
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Inches(6.1))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(cap, title, size=Pt(8.5), color=RGBColor(0x5B, 0x65, 0x73))


def _write_summary_csv(out_dir, results):
    summary_csv = os.path.join(out_dir, "批量指标汇总.csv")
    with open(summary_csv, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["run_number", "case_name", "steer_input_type", "duration_s",
                    "v_start_kmh", "v_end_kmh", "v_max_kmh", "max_abs_beta_deg",
                    "max_abs_yaw_degps", "max_abs_delta1_deg", "max_abs_delta3_deg"])
        for run_no, case, m, _fig_dir in results:
            w.writerow([run_no, case.get("case_name", ""), case.get("steer_input_type", ""),
                        _fmt(m["duration_s"]), _fmt(m["v_start_kmh"]), _fmt(m["v_end_kmh"]),
                        _fmt(m["v_max_kmh"]), _fmt(m["max_abs_beta_deg"]),
                        _fmt(m["max_abs_yaw_degps"]), _fmt(m["max_abs_delta1_deg"]),
                        _fmt(m["max_abs_delta3_deg"])])
    return summary_csv


def _load_results(dirs, out_dir):
    results = []
    for idx, run_dir in enumerate(dirs, start=1):
        run_no = os.path.basename(run_dir)
        data = load_run(run_dir, run_no)
        if data is None:
            print("SKIP_MISSING_CSV:", run_dir)
            continue
        io, py, case = data
        rows = build_table(io, py, case)
        if not rows:
            print("SKIP_EMPTY_SIGNAL:", run_dir)
            continue
        m = metrics(rows, case)
        write_metrics_csv(run_dir, run_no, m)
        safe_name = "%02d_%s_%s" % (idx, case.get("case_name", "case"), run_no)
        safe_name = "".join(ch if ch not in r'\/:*?"<>|' else "_" for ch in safe_name)
        fig_dir = os.path.join(out_dir, "figures", safe_name)
        os.makedirs(fig_dir, exist_ok=True)
        plot_figs(rows, fig_dir)
        results.append((run_no, case, m, fig_dir))
    return results


def main():
    ap = argparse.ArgumentParser(description="阶段二批量仿真总报告")
    ap.add_argument("--batch-marker", default=None, help="运行目录列表文件（每行一个目录）")
    ap.add_argument("--report-root", default=None, help="兼容旧参数；未指定 output-dir 时使用")
    ap.add_argument("--output-dir", default=None, help="批次统一输出目录")
    args = ap.parse_args()

    marker = args.batch_marker or os.environ.get("P2_BATCH_MARKER")
    if not marker:
        sys.exit("请指定 --batch-marker 或环境变量 P2_BATCH_MARKER")
    with open(marker, "r", encoding="utf-8-sig") as f:
        dirs = [ln.strip() for ln in f if ln.strip()]
    if not dirs:
        sys.exit("标记文件中没有运行目录")

    out_dir = args.output_dir or os.environ.get("P2_BATCH_OUTPUT_DIR")
    if not out_dir:
        report_root = args.report_root or os.environ.get("P2_REPORT_ROOT")
        if not report_root:
            sys.exit("请指定 --output-dir 或 --report-root")
        out_dir = os.path.join(report_root, "batch_summary")
    os.makedirs(out_dir, exist_ok=True)

    results = _load_results(dirs, out_dir)
    if not results:
        sys.exit("没有可汇总的运行")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(1.8)

    _add_title(doc, "阶段二批量仿真测试报告",
               "共 %d 个工况 | 输出目录：%s" % (len(results), out_dir))
    _add_summary_table(doc, results)

    add_heading(doc, "2. 分工况结果")
    for idx, (run_no, case, m, fig_dir) in enumerate(results, start=1):
        add_heading(doc, "2.%d %s（%s）" % (idx, case.get("case_name", ""), run_no))
        _add_case_table(doc, case, m)
        _add_case_figures(doc, fig_dir)

    summary_csv = _write_summary_csv(out_dir, results)
    docx_path = os.path.join(out_dir, "阶段二批量仿真测试报告.docx")
    doc.save(docx_path)

    print("BATCH_SUMMARY_OK", flush=True)


if __name__ == "__main__":
    main()
